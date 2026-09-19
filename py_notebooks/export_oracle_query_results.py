from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import tempfile
from pathlib import Path
from typing import Any, Iterable


MAX_EXCEL_ROWS = 1_048_576
FORMULA_PREFIXES = ("=", "+", "-", "@")


def load_env(path: Path) -> dict[str, str]:
   values: dict[str, str] = {}
   if not path.exists():
      return values
   for raw_line in path.read_text(encoding="utf-8").splitlines():
      line = raw_line.strip()
      if not line or line.startswith("#"):
         continue
      if line.startswith("export "):
         line = line[7:].lstrip()
      if "=" not in line:
         continue
      key, value = line.split("=", 1)
      value = value.strip()
      if len(value) >= 2 and value[0] == value[-1] and value[0] in {"'", '"'}:
         value = value[1:-1]
      values[key.strip()] = value
   return values


def make_dsn(config: dict[str, str]) -> str:
   dsn = config.get("ORACLE_DSN")
   if dsn:
      return dsn
   host = config.get("ORACLE_HOST")
   port = config.get("ORACLE_PORT", "1521")
   sid = config.get("ORACLE_SID")
   service_name = config.get("ORACLE_SERVICE_NAME")
   if not host or not (sid or service_name):
      raise ValueError(
         "Define ORACLE_DSN o bien ORACLE_HOST con ORACLE_SID/ORACLE_SERVICE_NAME."
      )
   import oracledb
   return oracledb.makedsn(host=host, port=int(port), sid=sid, service_name=service_name)


def load_manifest(path: Path) -> dict[str, Any]:
   try:
      manifest = json.loads(path.read_text(encoding="utf-8"))
   except json.JSONDecodeError as error:
      raise ValueError(f"El manifiesto no es JSON valido: {error}") from error
   if not isinstance(manifest, dict):
      raise ValueError("El manifiesto debe ser un objeto JSON.")
   for field in ("output", "format", "queries"):
      if field not in manifest:
         raise ValueError(f"Falta el campo obligatorio del manifiesto: {field}")
   if manifest["format"] not in {"xlsx", "csv"}:
      raise ValueError("format debe ser xlsx o csv.")
   if not isinstance(manifest["queries"], list) or not manifest["queries"]:
      raise ValueError("queries debe contener al menos una consulta.")
   for index, query in enumerate(manifest["queries"], start=1):
      if not isinstance(query, dict) or not isinstance(query.get("sql"), str):
         raise ValueError(f"La consulta {index} debe incluir sql como texto.")
      if not isinstance(query.get("binds", {}), dict):
         raise ValueError(f"Los binds de la consulta {index} deben ser un objeto.")
   return manifest


def replace_sql_literals(sql: str) -> str:
   output: list[str] = []
   index = 0
   while index < len(sql):
      character = sql[index]
      if character != "'":
         output.append(character)
         index += 1
         continue
      output.append(" ")
      index += 1
      while index < len(sql):
         if sql[index] == "'":
            if index + 1 < len(sql) and sql[index + 1] == "'":
               output.append("  ")
               index += 2
               continue
            index += 1
            break
         output.append(" ")
         index += 1
   return "".join(output)


def validate_select(sql: str) -> str:
   if not sql.strip():
      raise ValueError("La consulta no puede estar vacia.")
   outside_literals = replace_sql_literals(sql)
   if any(marker in outside_literals for marker in ("--", "/*", "*/", ";", "@")):
      raise ValueError("La consulta contiene comentarios, multiples sentencias o un enlace no permitido.")
   normalized = re.sub(r"\s+", " ", outside_literals).strip().upper()
   if not re.match(r"^(SELECT|WITH)\b", normalized):
      raise ValueError("Cada consulta debe iniciar con SELECT o WITH.")
   forbidden = (
      r"\b(INSERT|UPDATE|DELETE|MERGE|TRUNCATE|DROP|CREATE|ALTER|GRANT|REVOKE|"
      r"EXECUTE|IMMEDIATE|COMMIT|ROLLBACK|CALL|BEGIN|DECLARE)\b"
   )
   if re.search(forbidden, normalized) or re.search(r"\bFOR\s+UPDATE\b", normalized):
      raise ValueError("La consulta contiene una operacion o bloqueo no permitido.")
   if re.search(r"(?:\bSELECT\s+|,)\s*(?:[A-Z0-9_$#]+\.)*\*\s*(?:,|\bFROM\b|$)", normalized):
      raise ValueError("La exportacion requiere columnas explicitas; no se permite SELECT *.")
   return sql.strip()


def confirmation_hash(manifest: dict[str, Any]) -> str:
   payload = {
      "output": manifest["output"],
      "format": manifest["format"],
      "sheet_mode": manifest.get("sheet_mode", "one_per_result"),
      "queries": [
         {
            "name": query.get("name", f"resultado_{index}"),
            "sql": validate_select(query["sql"]),
            "binds": query.get("binds", {}),
         }
         for index, query in enumerate(manifest["queries"], start=1)
      ],
   }
   encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
   return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def parse_headers(description: Iterable[Any]) -> list[str]:
   headers: list[str] = []
   seen: dict[str, int] = {}
   for item in description:
      base = str(item[0]).strip() or "COLUMN"
      count = seen.get(base, 0) + 1
      seen[base] = count
      headers.append(base if count == 1 else f"{base}_{count}")
   return headers


def safe_cell(value: Any) -> Any:
   if hasattr(value, "read") and callable(value.read):
      value = value.read()
   if isinstance(value, (bytes, bytearray, memoryview)):
      raise ValueError("El resultado contiene un tipo binario no exportable.")
   if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
      return "'" + value
   return value


def write_result_csv(cursor: Any, temporary_path: Path) -> tuple[list[str], int]:
   headers = parse_headers(cursor.description or [])
   row_count = 0
   with temporary_path.open("w", encoding="utf-8", newline="") as stream:
      writer = csv.writer(stream)
      writer.writerow(headers)
      while True:
         rows = cursor.fetchmany(2_000)
         if not rows:
            break
         for row in rows:
            writer.writerow([safe_cell(value) for value in row])
            row_count += 1
   return headers, row_count


def safe_sheet_title(value: str, used: set[str]) -> str:
   title = re.sub(r"[\\/*?:\[\]]", "_", value).strip() or "Resultado"
   title = title[:31]
   candidate = title
   suffix = 2
   while candidate in used:
      suffix_text = f"_{suffix}"
      candidate = f"{title[:31 - len(suffix_text)]}{suffix_text}"
      suffix += 1
   used.add(candidate)
   return candidate


def safe_table_name(value: str, index: int) -> str:
   name = re.sub(r"[^A-Za-z0-9_]", "_", value)
   name = re.sub(r"^[^A-Za-z_]", "", name)[:200] or f"Resultado_{index}"
   if name[0].isdigit():
      name = f"R_{name}"
   return f"T_{name}_{index}"


def append_csv_to_sheet(worksheet: Any, path: Path, table_name: str, start_row: int) -> int:
   from openpyxl.worksheet.table import Table, TableStyleInfo
   with path.open("r", encoding="utf-8", newline="") as stream:
      reader = csv.reader(stream)
      first_row = start_row
      last_row = start_row - 1
      last_column = 0
      for row in reader:
         last_row += 1
         last_column = max(last_column, len(row))
         worksheet.append([safe_cell(value) for value in row])
   if last_row >= first_row and last_column:
      from openpyxl.utils import get_column_letter
      reference = f"A{first_row}:{get_column_letter(last_column)}{last_row}"
      table = Table(displayName=table_name, ref=reference)
      table.tableStyleInfo = TableStyleInfo(
         name="TableStyleMedium2",
         showFirstColumn=False,
         showLastColumn=False,
         showRowStripes=True,
         showColumnStripes=False,
      )
      worksheet.add_table(table)
   return last_row


def write_xlsx(results: list[dict[str, Any]], output_path: Path, sheet_mode: str) -> None:
   from openpyxl import Workbook
   workbook = Workbook()
   workbook.remove(workbook.active)
   used_titles: set[str] = set()
   if sheet_mode == "same_sheet":
      worksheet = workbook.create_sheet(safe_sheet_title("Resultados", used_titles))
      row = 1
      for index, result in enumerate(results, start=1):
         row = append_csv_to_sheet(
            worksheet, result["temporary_path"], safe_table_name(result["name"], index), row
         ) + 2
   else:
      for index, result in enumerate(results, start=1):
         worksheet = workbook.create_sheet(safe_sheet_title(result["name"], used_titles))
         append_csv_to_sheet(
            worksheet, result["temporary_path"], safe_table_name(result["name"], index), 1
         )
         worksheet.freeze_panes = "A2"
   output_path.parent.mkdir(parents=True, exist_ok=True)
   workbook.save(output_path)


def output_csv_paths(output_path: Path, results: list[dict[str, Any]]) -> list[Path]:
   if len(results) == 1:
      return [output_path]
   paths: list[Path] = []
   for index, result in enumerate(results, start=1):
      slug = re.sub(r"[^A-Za-z0-9_-]+", "_", result["name"]).strip("_") or f"resultado_{index}"
      paths.append(output_path.with_name(f"{output_path.stem}_{slug}.csv"))
   return paths


def connect(config: dict[str, str]) -> Any:
   import oracledb
   user = config.get("ORACLE_USER")
   password = config.get("ORACLE_PASSWORD") or config.get("ORACLE_PWD")
   if not user or not password:
      raise ValueError("Define ORACLE_USER y ORACLE_PASSWORD en el archivo .env.")
   connection_args: dict[str, Any] = {"user": user, "password": password, "dsn": make_dsn(config)}
   if config.get("ORACLE_CONFIG_DIR"):
      connection_args["config_dir"] = config["ORACLE_CONFIG_DIR"]
   return oracledb.connect(**connection_args)


def export(manifest: dict[str, Any], env_path: Path) -> list[Path]:
   actual_hash = confirmation_hash(manifest)
   if manifest.get("confirmed_query_hash") != actual_hash:
      raise ValueError(
         "La confirmacion no coincide con el manifiesto actual. "
         f"Hash esperado por el ejecutor: {actual_hash}"
      )
   output_path = Path(manifest["output"]).expanduser()
   requested_format = manifest["format"]
   if requested_format == "xlsx" and output_path.suffix.lower() != ".xlsx":
      raise ValueError("Una salida XLSX debe terminar en .xlsx.")
   if requested_format == "csv" and output_path.suffix.lower() not in {".csv", ".xlsx"}:
      raise ValueError("Una salida CSV debe terminar en .csv o usar .xlsx como nombre base.")
   sheet_mode = manifest.get("sheet_mode", "one_per_result")
   if sheet_mode not in {"same_sheet", "one_per_result"}:
      raise ValueError("sheet_mode debe ser same_sheet o one_per_result.")
   if output_path.exists():
      raise FileExistsError(f"El archivo de salida ya existe: {output_path}")

   output_path.parent.mkdir(parents=True, exist_ok=True)
   temporary_paths: list[Path] = []
   results: list[dict[str, Any]] = []
   connection = None
   try:
      connection = connect(load_env(env_path))
      for index, query in enumerate(manifest["queries"], start=1):
         name = str(query.get("name") or f"resultado_{index}")
         descriptor, temporary_name = tempfile.mkstemp(
            prefix="oracle_export_", suffix=".csv", dir=output_path.parent
         )
         os.close(descriptor)
         temporary = Path(temporary_name)
         temporary_paths.append(temporary)
         cursor = connection.cursor()
         try:
            cursor.execute(validate_select(query["sql"]), query.get("binds", {}))
            headers, row_count = write_result_csv(cursor, temporary)
         finally:
            cursor.close()
         results.append(
            {"name": name, "temporary_path": temporary, "headers": headers, "row_count": row_count}
         )
   finally:
      if connection is not None:
         connection.close()

   too_large = any(result["row_count"] >= MAX_EXCEL_ROWS for result in results)
   effective_format = "csv" if too_large else requested_format
   if effective_format == "xlsx":
      write_xlsx(results, output_path, sheet_mode)
      created = [output_path]
   else:
      csv_output = output_path.with_suffix(".csv") if output_path.suffix.lower() == ".xlsx" else output_path
      created = output_csv_paths(csv_output, results)
      for result, destination in zip(results, created):
         if destination.exists():
            raise FileExistsError(f"El archivo de salida ya existe: {destination}")
         os.replace(result["temporary_path"], destination)
   for temporary in temporary_paths:
      if temporary.exists():
         temporary.unlink()
   if too_large and requested_format == "xlsx":
      print("El resultado supera el limite de XLSX; se forzo salida CSV.")
   summary = "; ".join(
      f"{result['name']}: {result['row_count']} filas, {len(result['headers'])} columnas"
      for result in results
   )
   print(f"Exportacion completada ({effective_format}): {', '.join(map(str, created))}")
   print(summary)
   return created


def parse_args() -> argparse.Namespace:
   parser = argparse.ArgumentParser(description="Exporta SELECT Oracle a XLSX o CSV.")
   parser.add_argument("--manifest", required=True, help="Manifiesto JSON con consultas, binds y salida.")
   parser.add_argument("--env-file", default=".env", help="Archivo de configuracion Oracle.")
   parser.add_argument(
      "--print-query-hash", action="store_true", help="Muestra el hash sin conectarse a Oracle."
   )
   return parser.parse_args()


def main() -> None:
   args = parse_args()
   manifest = load_manifest(Path(args.manifest).expanduser())
   if args.print_query_hash:
      print(confirmation_hash(manifest))
      return
   export(manifest, Path(args.env_file).expanduser())


if __name__ == "__main__":
   main()




