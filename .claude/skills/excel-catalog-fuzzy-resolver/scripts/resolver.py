import argparse
import csv
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from rapidfuzz import fuzz

SCRIPT_DIR = Path(__file__).parent
DEFAULT_ALIASES = SCRIPT_DIR / "aliases_es.json"

ID_REGEX = re.compile(r"(?:^ID[_.])|(?:_ID[_.])|(?:_ID$)|(?:ID$)", re.IGNORECASE)

PESOS_SCORERS = {
    "token_set_ratio": 0.40,
    "token_sort_ratio": 0.25,
    "partial_ratio": 0.20,
    "wratio": 0.15,
}

UMBRAL_SCORE = 82
UMBRAL_INCIERTO = 95
UMBRAL_MAYORIA = 0.80
UMBRAL_MINIMO_BOOST = 70


def fail(msg, code=1):
    print(f"ERROR: {msg}", file=sys.stderr)
    sys.exit(code)


def normalizar(texto):
    if texto is None:
        return ""
    s = str(texto)
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = s.lower()
    for ch in [".", ":", "(", ")", "/", "-", "_", ",", "\n", "\t"]:
        s = s.replace(ch, " ")
    s = " ".join(s.split()).strip()
    return s


def parse_col(col_str):
    col_str = str(col_str).strip()
    if not col_str:
        return None
    if col_str.isdigit():
        return int(col_str)
    try:
        return column_index_from_string(col_str.upper())
    except Exception:
        return None


def parse_category_ranges(arg):
    if not arg:
        return {}
    out = {}
    for chunk in arg.split(";"):
        chunk = chunk.strip()
        if not chunk:
            continue
        if "=" not in chunk:
            continue
        key, rango = chunk.split("=", 1)
        key = key.strip()
        ini, fin = rango.split("-", 1)
        out[key] = (int(ini), int(fin))
    return out


def cargar_catalogo(source_path, source_sheet, key_col, result_col):
    wb = load_workbook(source_path, data_only=True)
    if source_sheet:
        if source_sheet not in wb.sheetnames:
            fail(f"Hoja '{source_sheet}' no existe en {source_path}. Hojas disponibles: {wb.sheetnames}")
        ws = wb[source_sheet]
    else:
        ws = wb.active
    catalogo = []
    for r in range(1, ws.max_row + 1):
        desc = ws.cell(row=r, column=key_col).value
        cid = ws.cell(row=r, column=result_col).value
        if desc is None or cid is None:
            continue
        catalogo.append((normalizar(desc), cid, desc))
    return catalogo, ws.max_row


def filtrar_por_categoria(catalogo, rango):
    if not rango:
        return list(catalogo)
    ini, fin = rango
    ini_idx = ini - 1
    fin_idx = fin
    return [c for i, c in enumerate(catalogo) if ini_idx <= i < fin_idx]


def score_ensemble(valor_norm, candidato_norm):
    return (
        PESOS_SCORERS["token_set_ratio"] * fuzz.token_set_ratio(valor_norm, candidato_norm)
        + PESOS_SCORERS["token_sort_ratio"] * fuzz.token_sort_ratio(valor_norm, candidato_norm)
        + PESOS_SCORERS["partial_ratio"] * fuzz.partial_ratio(valor_norm, candidato_norm)
        + PESOS_SCORERS["wratio"] * fuzz.WRatio(valor_norm, candidato_norm)
    )


def resolver(valor_norm, candidatos, alias_map):
    if not valor_norm:
        return None, 0.0, None, "vacio"
    if alias_map and valor_norm in alias_map:
        target = alias_map[valor_norm]
        for c in candidatos:
            if c[0] == target:
                return c[1], 100.0, c[2], "alias_exacto"
        for c in candidatos:
            if target in c[0] or c[0] in target:
                return c[1], 95.0, c[2], "alias_parcial"
    for c in candidatos:
        if c[0] == valor_norm:
            return c[1], 100.0, c[2], "exacto"
    mejor = None
    mejor_score = -1.0
    for c in candidatos:
        s = score_ensemble(valor_norm, c[0])
        if s > mejor_score:
            mejor_score = s
            mejor = c
    if mejor is None:
        return None, 0.0, None, "sin_candidatos"
    return mejor[1], mejor_score, mejor[2], "ensemble"


def detectar_columnas_id(ws, header_row, id_regex, skip_cols=None):
    skip_idx = set()
    skip_names = set()
    if skip_cols:
        for token in skip_cols.split(","):
            token = token.strip()
            if not token:
                continue
            idx = parse_col(token)
            if idx:
                skip_idx.add(idx)
            else:
                skip_names.add(token.upper())
    headers = {}
    for c in range(1, ws.max_column + 1):
        h = ws.cell(row=header_row, column=c).value
        headers[c] = h
    detectadas = []
    omitidas = []
    for c, h in headers.items():
        if h is None or str(h).strip() == "":
            continue
        if id_regex.search(str(h)):
            if c in skip_idx or str(h).upper() in skip_names:
                omitidas.append((c, h, "excluida por --skip-columns"))
                continue
            desc_col = c + 1
            if desc_col > ws.max_column:
                omitidas.append((c, h, "adyacente derecha fuera de rango"))
                continue
            desc_header = ws.cell(row=header_row, column=desc_col).value
            if desc_header is None or str(desc_header).strip() == "":
                omitidas.append((c, h, f"adyacente col {get_column_letter(desc_col)} sin header"))
                continue
            detectadas.append((c, get_column_letter(c), h, desc_col, get_column_letter(desc_col), desc_header))
    return detectadas, omitidas, headers


def confirmar(detectadas, omitidas, no_confirm):
    print("\n=== Columnas ID detectadas ===")
    for c, col_letter, header, desc_col, desc_letter, desc_header in detectadas:
        print(f"  {col_letter} '{header}'  ->  adyacente {desc_letter} '{desc_header}'")
    if omitidas:
        print("\n=== Columnas ID omitidas ===")
        for c, h, razon in omitidas:
            print(f"  col {get_column_letter(c)} '{h}': {razon}")
    if no_confirm:
        print("\n(--no-confirm activo, no se solicita confirmacion)")
        return True
    print("\n¿Continuar con la escritura? (s/N): ", end="", flush=True)
    resp = input().strip().lower()
    return resp in ("s", "si", "y", "yes")


def main():
    parser = argparse.ArgumentParser(
        description="Pobla columnas ID_ en un Excel target mediante busqueda fuzzy contra un catalogo source.",
    )
    parser.add_argument("--target", required=True, help="Ruta absoluta del .xlsx destino")
    parser.add_argument("--target-header-row", type=int, required=True, help="Fila 1-based de encabezados del target")
    parser.add_argument("--target-sheet", help="Nombre de la hoja del target (default: primera)")
    parser.add_argument("--source", required=True, help="Ruta absoluta del .xlsx catalogo")
    parser.add_argument("--source-key-col", required=True, help="Columna clave del catalogo (letra o numero)")
    parser.add_argument("--source-result-col", required=True, help="Columna ID del catalogo (letra o numero)")
    parser.add_argument("--source-sheet", help="Nombre de la hoja del catalogo (default: primera)")
    parser.add_argument("--umbral", type=int, default=UMBRAL_SCORE, help=f"Umbral de score (default {UMBRAL_SCORE})")
    parser.add_argument("--category-ranges", help="Pre-filtro por rangos: 'COL=ini-fin;COL2=ini2-fin2'")
    parser.add_argument("--aliases-file", help="Archivo JSON con aliases adicionales")
    parser.add_argument("--output-suffix", default="_resuelto", help="Sufijo del archivo de salida (default _resuelto)")
    parser.add_argument("--no-confirm", action="store_true", help="Saltar confirmacion de columnas detectadas")
    parser.add_argument("--id-regex", help="Regex personalizada para detectar columnas ID")
    parser.add_argument("--skip-columns", help="Columnas a excluir (separadas por coma; por letra, numero o nombre de header). Ej: 'B,E' o 'ZO_ID_ZONA,PF_ID_FAMILIA'")

    args = parser.parse_args()

    target_path = Path(args.target)
    source_path = Path(args.source)

    if not target_path.is_file():
        fail(f"El archivo target no existe: {args.target}")
    if not source_path.is_file():
        fail(f"El archivo source no existe: {args.source}")
    if args.target_header_row < 1:
        fail(f"--target-header-row debe ser >= 1, recibido {args.target_header_row}")

    source_key_idx = parse_col(args.source_key_col)
    source_result_idx = parse_col(args.source_result_col)
    if not source_key_idx:
        fail(f"--source-key-col invalido: {args.source_key_col}")
    if not source_result_idx:
        fail(f"--source-result-col invalido: {args.source_result_col}")

    id_regex = re.compile(args.id_regex, re.IGNORECASE) if args.id_regex else ID_REGEX
    category_ranges = parse_category_ranges(args.category_ranges)

    aliases = {}
    if DEFAULT_ALIASES.is_file():
        with DEFAULT_ALIASES.open(encoding="utf-8") as f:
            aliases = json.load(f)
    if args.aliases_file:
        ap = Path(args.aliases_file)
        if not ap.is_file():
            fail(f"Archivo de aliases no existe: {args.aliases_file}")
        with ap.open(encoding="utf-8") as f:
            extra = json.load(f)
        for k, v in extra.items():
            aliases.setdefault(k, {}).update(v)

    print(f"Cargando catalogo: {source_path}")
    catalogo_full, _ = cargar_catalogo(source_path, args.source_sheet, source_key_idx, source_result_idx)
    print(f"  {len(catalogo_full)} entradas en catalogo")

    print(f"\nCargando target: {target_path}")
    wb = load_workbook(target_path)
    if args.target_sheet:
        if args.target_sheet not in wb.sheetnames:
            fail(f"Hoja '{args.target_sheet}' no existe en {target_path}. Hojas: {wb.sheetnames}")
        ws = wb[args.target_sheet]
    else:
        ws = wb.active
    print(f"  Hoja: {ws.title}, filas: {ws.max_row}, cols: {ws.max_column}")

    print(f"\nDetectando columnas ID en fila {args.target_header_row} con regex: {id_regex.pattern}")
    detectadas, omitidas, headers = detectar_columnas_id(ws, args.target_header_row, id_regex, args.skip_columns)
    if not detectadas:
        fail("Ninguna columna ID detectada con la regex actual. Ajustala o usa --id-regex.")
    if not confirmar(detectadas, omitidas, args.no_confirm):
        print("Cancelado por el usuario.")
        sys.exit(0)

    print("\nPasada 1: scoring para todas las celdas...")
    celdas = []
    for fila in range(args.target_header_row + 1, ws.max_row + 1):
        for c_id, col_letter, header_id, desc_col, desc_letter, desc_header in detectadas:
            val_desc = ws.cell(row=fila, column=desc_col).value
            if val_desc is None or str(val_desc).strip() == "":
                continue
            val_norm = normalizar(val_desc)
            rango = category_ranges.get(header_id)
            candidatos = filtrar_por_categoria(catalogo_full, rango)
            alias_map = aliases.get(header_id, {})
            cid, score, desc_orig, metodo = resolver(val_norm, candidatos, alias_map)
            celdas.append({
                "fila": fila,
                "c_id": c_id,
                "col_id_name": header_id,
                "val_desc": val_desc,
                "val_norm": val_norm,
                "candidato_id": cid,
                "score": score,
                "desc_orig": desc_orig,
                "metodo": metodo,
            })
    print(f"  celdas a evaluar: {len(celdas)}")

    print("\nPasada 2: boost por mayoria historica...")
    frecuencias = defaultdict(Counter)
    for c in celdas:
        if c["candidato_id"] is not None:
            key = (c["col_id_name"], c["val_norm"])
            frecuencias[key][c["candidato_id"]] += 1
    boosts = 0
    for c in celdas:
        key = (c["col_id_name"], c["val_norm"])
        total = sum(frecuencias[key].values())
        if total == 0:
            continue
        ganador, freq = frecuencias[key].most_common(1)[0]
        ratio = freq / total
        if (
            c["score"] < UMBRAL_INCIERTO
            and ratio >= UMBRAL_MAYORIA
            and c["score"] >= UMBRAL_MINIMO_BOOST
            and ganador == c["candidato_id"]
        ):
            c["score"] = max(c["score"], 95.0)
            c["metodo"] = "mayoria_boost"
            boosts += 1
    print(f"  boosts aplicados: {boosts}")

    print("\nEscribiendo resultados...")
    cobertura = {header_id: {"match": 0, "total": 0} for _, _, header_id, _, _, _ in detectadas}
    no_match = []
    ambiguas = []
    for c in celdas:
        cobertura[c["col_id_name"]]["total"] += 1
        if c["candidato_id"] is not None and c["score"] >= args.umbral:
            ws.cell(row=c["fila"], column=c["c_id"]).value = c["candidato_id"]
            cobertura[c["col_id_name"]]["match"] += 1
            if c["score"] < 90:
                ambiguas.append({
                    "fila": c["fila"],
                    "columna_id": c["col_id_name"],
                    "valor_origen": c["val_desc"],
                    "id_asignado": c["candidato_id"],
                    "desc_catalogo": c["desc_orig"],
                    "score": round(c["score"], 2),
                    "metodo": c["metodo"],
                })
        else:
            no_match.append({
                "fila": c["fila"],
                "columna_id": c["col_id_name"],
                "valor_origen": c["val_desc"],
                "mejor_candidato": c["desc_orig"],
                "score": round(c["score"], 2),
                "metodo": c["metodo"],
            })

    out_path = target_path.with_name(target_path.stem + args.output_suffix + target_path.suffix)
    print(f"\nGuardando: {out_path}")
    wb.save(out_path)

    rep_dir = target_path.parent
    rep_cobertura = rep_dir / "reporte_cobertura.txt"
    rep_nomatch = rep_dir / "reporte_no_match.csv"
    rep_ambig = rep_dir / "reporte_ambiguedades.csv"
    rep_deteccion = rep_dir / "reporte_deteccion.txt"

    lineas = ["REPORTE DE COBERTURA - excel-catalog-fuzzy-resolver", "=" * 60]
    for col, stats in cobertura.items():
        total = stats["total"]
        match = stats["match"]
        pct = (match / total * 100) if total else 0
        lineas.append(f"  {col:30s}  match={match:5d}/{total:5d}  ({pct:5.1f}%)")
    total_match = sum(s["match"] for s in cobertura.values())
    total_total = sum(s["total"] for s in cobertura.values())
    pct_global = (total_match / total_total * 100) if total_total else 0
    lineas.append(f"\n  GLOBAL                              match={total_match:5d}/{total_total:5d}  ({pct_global:5.1f}%)")
    lineas.append(f"  Boosts por mayoria: {boosts}")
    lineas.append(f"  Celdas ambiguas (score<90): {len(ambiguas)}")
    rep_cobertura.write_text("\n".join(lineas), encoding="utf-8")

    with rep_nomatch.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["fila", "columna_id", "valor_origen", "mejor_candidato", "score", "metodo"])
        w.writeheader()
        w.writerows(no_match)

    with rep_ambig.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["fila", "columna_id", "valor_origen", "id_asignado", "desc_catalogo", "score", "metodo"])
        w.writeheader()
        w.writerows(ambiguas)

    det_lineas = ["REPORTE DE DETECCION DE COLUMNAS ID", "=" * 60, f"Regex: {id_regex.pattern}", f"Fila de encabezados: {args.target_header_row}", ""]
    det_lineas.append("Detectadas y procesadas:")
    for c, col_letter, header, desc_col, desc_letter, desc_header in detectadas:
        det_lineas.append(f"  {col_letter} '{header}' -> adyacente {desc_letter} '{desc_header}'")
    if omitidas:
        det_lineas.append("")
        det_lineas.append("Omitidas:")
        for c, h, razon in omitidas:
            det_lineas.append(f"  {get_column_letter(c)} '{h}': {razon}")
    rep_deteccion.write_text("\n".join(det_lineas), encoding="utf-8")

    print("\n" + "\n".join(lineas))
    print(f"\nReportes:")
    print(f"  - {rep_cobertura}")
    print(f"  - {rep_nomatch}")
    print(f"  - {rep_ambig}")
    print(f"  - {rep_deteccion}")
    print(f"\nArchivo de salida: {out_path}")


if __name__ == "__main__":
    main()
